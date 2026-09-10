"""
Persistent per-operator global Knowledge Base + inline document preview.

Covers:
  - docsearch: global-KB ingest / search isolation by user_id, the
    search_all chat+global merge, and delete.
  - doc_preview: server-side offline rendering for each supported format.
  - the /api/kb/* and /api/preview/* HTTP endpoints, including the
    path-traversal guards.

All offline: ChromaDB is local/on-disk and the preview path uses the same
python-docx / openpyxl already pinned for filegen. Every test uses a unique
user_id and cleans up its own chunks so the shared dev Chroma store isn't
polluted.
"""
import base64
import io
import uuid

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.tools import docsearch, doc_preview
from app.storage.config import FILES_DIR


client = TestClient(app)


def _uid() -> str:
    return f"kbtest-{uuid.uuid4().hex[:12]}"


def _auth(user_id: str | None = None) -> tuple[str, dict]:
    """Create a real server-side account and return (user_id, auth headers).

    The /api/kb/* endpoints take identity from the bearer token now, never
    from a caller-supplied user_id — so every KB test has to authenticate
    exactly the way the browser does.
    """
    user_id = user_id or _uid()
    r = client.post("/api/auth/signup", json={"user_id": user_id, "password": "pw-test-1234"})
    assert r.status_code == 200, r.text
    return user_id, {"Authorization": f"Bearer {r.json()['token']}"}


def _txt_b64(text: str) -> str:
    return base64.b64encode(text.encode("utf-8")).decode("ascii")


def _docx_bytes(paragraphs) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# docsearch: global KB
# --------------------------------------------------------------------------

def test_global_kb_ingest_and_search_is_user_scoped():
    user_a, user_b = _uid(), _uid()
    try:
        docsearch.ingest_global_document(
            user_id=user_a, filename="alpha.txt",
            pages=[(1, "The turbine bearing clearance limit is 0.15 mm per spec TB-9.")],
        )
        docsearch.ingest_global_document(
            user_id=user_b, filename="beta.txt",
            pages=[(1, "The safety valve set pressure is 12.5 barg per document SV-4.")],
        )

        hits_a = docsearch.search_global_kb("bearing clearance limit", user_a, top_k=3)
        assert hits_a and hits_a[0]["source"] == "alpha.txt"
        assert all(h.get("origin") == "global-kb" for h in hits_a)

        # user_a's query must never surface user_b's document.
        assert all(h["source"] != "beta.txt" for h in hits_a)
        hits_b = docsearch.search_global_kb("bearing clearance limit", user_b, top_k=3)
        assert all(h["source"] != "alpha.txt" for h in hits_b)
    finally:
        for u, d in ((user_a, "alpha.txt"), (user_b, "beta.txt")):
            for doc in docsearch.list_global_documents(u):
                docsearch.delete_global_document(u, doc["document_id"])


def test_search_all_merges_chat_and_global():
    user_id = _uid()
    chat_id = f"chat-{uuid.uuid4().hex[:8]}"
    try:
        docsearch.ingest_global_document(
            user_id=user_id, filename="global-policy.txt",
            pages=[(1, "Global rule: permits expire after 7 days.")],
        )
        docsearch.ingest_chat_document(
            chat_id=chat_id, filename="chat-note.txt",
            pages=[(1, "This chat's note: the permit was issued on Monday.")],
        )
        merged = docsearch.search_all("permit", top_k=5, chat_id=chat_id, user_id=user_id)
        sources = {m["source"] for m in merged}
        assert "global-policy.txt" in sources
        assert "chat-note.txt" in sources
    finally:
        for doc in docsearch.list_global_documents(user_id):
            docsearch.delete_global_document(user_id, doc["document_id"])
        try:
            docsearch.get_chat_collection().delete(where={"chat_id": chat_id})
        except Exception:
            pass


def test_global_kb_delete_removes_chunks():
    user_id = _uid()
    res = docsearch.ingest_global_document(
        user_id=user_id, filename="scratch.txt",
        pages=[(1, "disposable content about pumps and seals")],
    )
    assert res["chunks"] >= 1
    removed = docsearch.delete_global_document(user_id, res["document_id"])
    assert removed >= 1
    assert docsearch.list_global_documents(user_id) == []
    assert docsearch.search_global_kb("pumps and seals", user_id, top_k=3) == []


# --------------------------------------------------------------------------
# doc_preview
# --------------------------------------------------------------------------

def test_preview_pdf_mode_is_passthrough():
    assert doc_preview.render_preview(b"%PDF-1.4 fake", "report.pdf")["mode"] == "pdf"


def test_preview_docx_renders_html():
    raw = _docx_bytes(["Executive Summary", "The plant met its uptime target."])
    out = doc_preview.render_preview(raw, "summary.docx")
    assert out["mode"] == "html"
    assert "uptime target" in out["html"]
    assert "<script" not in out["html"].lower()


def test_preview_xlsx_renders_table_per_sheet():
    raw = _xlsx_bytes([["Name", "Value"], ["Flow", 42], ["Pressure", "7 barg"]])
    out = doc_preview.render_preview(raw, "data.xlsx")
    assert out["mode"] == "html"
    assert "<table" in out["html"]
    assert "7 barg" in out["html"]
    assert "Sheet1" in out["html"]


def test_preview_text_mode_and_escaping():
    out = doc_preview.render_preview(b"line one\n<b>not html</b>", "notes.txt")
    assert out["mode"] == "text"
    assert out["text"] == "line one\n<b>not html</b>"


def test_preview_unsupported_type():
    assert doc_preview.render_preview(b"\x00\x01", "archive.zip")["mode"] == "unsupported"


# --------------------------------------------------------------------------
# HTTP endpoints
# --------------------------------------------------------------------------

def test_kb_endpoints_full_lifecycle():
    user_id, hdrs = _auth()
    up = client.post("/api/kb/upload", headers=hdrs, json={
        "file_name": "handbook.txt",
        "file_base64": _txt_b64("Section 4: lockout-tagout must be verified by a second person."),
        "file_mime_type": "text/plain",
    })
    assert up.status_code == 200, up.text
    doc_id = up.json()["document_id"]
    assert up.json()["chunks"] >= 1

    listing = client.get("/api/kb/list", headers=hdrs)
    assert listing.status_code == 200
    docs = listing.json()["documents"]
    assert len(docs) == 1 and docs[0]["filename"] == "handbook.txt"

    prev = client.get(f"/api/kb/{doc_id}/preview", headers=hdrs)
    assert prev.status_code == 200
    body = prev.json()
    assert body["mode"] == "text"
    assert "lockout-tagout" in body["text"]
    assert body["raw_url"].startswith(f"/api/kb/{doc_id}/raw")

    raw = client.get(f"/api/kb/{doc_id}/raw", headers=hdrs)
    assert raw.status_code == 200
    assert b"lockout-tagout" in raw.content

    delr = client.request("DELETE", f"/api/kb/{doc_id}", headers=hdrs)
    assert delr.status_code == 200 and delr.json()["success"] is True
    assert client.get("/api/kb/list", headers=hdrs).json()["documents"] == []


def test_kb_upload_rejects_unsupported_and_images():
    _, hdrs = _auth()
    bad = client.post("/api/kb/upload", headers=hdrs, json={
        "file_name": "photo.png",
        "file_base64": _txt_b64("x"), "file_mime_type": "image/png",
    })
    assert bad.status_code == 400

    bad2 = client.post("/api/kb/upload", headers=hdrs, json={
        "file_name": "app.exe", "file_base64": _txt_b64("x"),
    })
    assert bad2.status_code == 400


def test_kb_path_traversal_is_rejected():
    """user_id can no longer be supplied by the caller at all (it comes from
    the token), which removes that traversal vector structurally. The
    document_id segment is still caller-controlled, so it stays guarded."""
    _, hdrs = _auth()
    r = client.get("/api/kb/%2e%2e/preview", headers=hdrs)
    assert r.status_code in (400, 404)
    r = client.get("/api/kb/%2e%2e/raw", headers=hdrs)
    assert r.status_code in (400, 404)


def test_kb_requires_authentication():
    """Every /api/kb/* route rejects an unauthenticated caller. Before the
    server-side session existed these all answered anyone on the LAN."""
    for method, path in [
        ("GET", "/api/kb/list"),
        ("GET", "/api/kb/some-doc/preview"),
        ("GET", "/api/kb/some-doc/raw"),
        ("DELETE", "/api/kb/some-doc"),
    ]:
        r = client.request(method, path)
        assert r.status_code == 401, f"{method} {path} -> {r.status_code}"
    r = client.post("/api/kb/upload", json={"file_name": "a.txt", "file_base64": _txt_b64("x")})
    assert r.status_code == 401


def test_kb_is_isolated_between_authenticated_users():
    """The actual vulnerability this replaced: user_id was a string the
    caller chose, so anyone could list, read, or delete another operator's
    documents just by naming them. Identity now comes from the token, so
    user B simply cannot address user A's files."""
    _, hdrs_a = _auth()
    _, hdrs_b = _auth()

    up = client.post("/api/kb/upload", headers=hdrs_a, json={
        "file_name": "confidential.txt",
        "file_base64": _txt_b64("turbine bearing clearance is 0.42mm"),
        "file_mime_type": "text/plain",
    })
    assert up.status_code == 200, up.text
    doc_id = up.json()["document_id"]

    # B cannot see A's document in their own listing...
    assert client.get("/api/kb/list", headers=hdrs_b).json()["documents"] == []
    # ...nor read it by naming A's document_id directly...
    assert client.get(f"/api/kb/{doc_id}/raw", headers=hdrs_b).status_code == 404
    assert client.get(f"/api/kb/{doc_id}/preview", headers=hdrs_b).status_code == 404
    # ...and a delete by B must not destroy A's file.
    client.request("DELETE", f"/api/kb/{doc_id}", headers=hdrs_b)
    assert client.get(f"/api/kb/{doc_id}/raw", headers=hdrs_a).status_code == 200

    client.request("DELETE", f"/api/kb/{doc_id}", headers=hdrs_a)


def test_preview_generated_file_in_files_dir():
    name = f"kbtest-{uuid.uuid4().hex[:8]}.txt"
    path = FILES_DIR / name
    path.write_text("Deliverable body: quarterly totals attached.", encoding="utf-8")
    try:
        r = client.get(f"/api/preview/generated/{name}")
        assert r.status_code == 200
        body = r.json()
        assert body["mode"] == "text"
        assert "quarterly totals" in body["text"]
        assert body["raw_url"] == f"/files/{name}"
    finally:
        path.unlink(missing_ok=True)

    assert client.get("/api/preview/generated/nope-missing.txt").status_code == 404
    assert client.get("/api/preview/generated/..%2fconfig.json").status_code in (400, 404)
